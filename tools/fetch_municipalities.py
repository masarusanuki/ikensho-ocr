# -*- coding: utf-8 -*-
"""総務省の「都道府県コード及び市区町村コード」から市区町村の辞書を作る。

住所欄（申請者の住所・医療機関所在地）の読み取りを、一覧と突き合わせて直すため。
誤字表（tools/build_word_fixes.py）の種にもなる。

    python3 tools/fetch_municipalities.py
    python3 tools/fetch_municipalities.py --from-file 000925835.xlsx

出典: 総務省 https://www.soumu.go.jp/denshijiti/code.html
      「都道府県コード及び市区町村コード」（Excel）

**取得できない環境では、手元の医療機関一覧（dict/hospitals/）の住所から
市区町村名を拾う**（--from-hospitals）。こちらは医療機関のある市区町村しか
入らないが、辞書としては十分に使える。
"""
import argparse
import glob
import json
import os
import re
import sys
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "dict", "municipalities.json")
PAGE = "https://www.soumu.go.jp/denshijiti/code.html"
UA = {"User-Agent": "ikensho-ocr/1.0 (+https://github.com/masarusanuki/ikensho-ocr)"}

# 市区町村名の末尾。ここで切ると「○○市」「○○郡××町」が取れる
SUFFIX = re.compile(r"^(.+?[市区町村])")
# 政令指定都市の区は「札幌市中央区」の形で入っている。市の部分も別に登録する
CITY_WARD = re.compile(r"^(.+?市)(.+区)$")


def find_xlsx_url():
    """一覧ページから Excel の場所を見つける。"""
    with urllib.request.urlopen(urllib.request.Request(PAGE, headers=UA), timeout=60) as r:
        html = r.read().decode("cp932", "replace")
    m = re.findall(r'href="(/main_content/\d+\.xlsx)"', html)
    if not m:
        raise SystemExit("Excel の場所が見つかりませんでした。--from-file で渡してください。")
    return "https://www.soumu.go.jp" + m[0]


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = []
    for name in wb.sheetnames:
        ws = wb[name]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or len(r) < 3:
                continue
            pref, city = r[1], r[2]
            if pref and city:
                rows.append((str(pref).strip(), str(city).strip()))
    return rows


def from_hospitals():
    """医療機関一覧の住所から市区町村名を拾う（取得できないときの代わり）。"""
    rows = []
    for path in sorted(glob.glob(os.path.join(ROOT, "dict", "hospitals", "*.json"))):
        if path.endswith("index.json"):
            continue
        with open(path, encoding="utf-8") as fp:
            data = json.load(fp)
        pref = data.get("pref", "")
        for e in data.get("entries", []):
            m = SUFFIX.match(str(e.get("address") or ""))
            if m:
                rows.append((pref, m.group(1)))
    return rows


def build(rows):
    """(都道府県, 市区町村) の並びを、重複の無い辞書にする。"""
    seen = {}
    for pref, city in rows:
        for name in [city] + ([CITY_WARD.match(city).group(1)]
                              if CITY_WARD.match(city) else []):
            if not name or not re.match(r"^[ぁ-んァ-ヶ一-龥ヶヵ]+$", name):
                continue
            key = (pref, name)
            if key not in seen:
                seen[key] = dict(name=name, pref=pref)
    return [seen[k] for k in sorted(seen)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--from-file", help="ダウンロード済みの Excel")
    ap.add_argument("--from-hospitals", action="store_true",
                    help="医療機関一覧の住所から拾う（取得できないとき）")
    ap.add_argument("--out", default=OUT)
    args = ap.parse_args()

    if args.from_hospitals:
        rows = from_hospitals()
        source = "dict/hospitals/（医療機関一覧の住所から）"
    else:
        path = args.from_file
        if not path:
            url = find_xlsx_url()
            print(f"取得します: {url}", file=sys.stderr)
            path = os.path.join(ROOT, "dict", "_municipalities.xlsx")
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA),
                                        timeout=120) as r, open(path, "wb") as fp:
                fp.write(r.read())
        rows = read_xlsx(path)
        source = "総務省 都道府県コード及び市区町村コード"
        if not args.from_file and os.path.exists(path):
            os.remove(path)

    entries = build(rows)
    payload = dict(label="市区町村", source=source,
                   note="住所欄の読み取りを突き合わせるための一覧。"
                        "tools/fetch_municipalities.py が作ります。",
                   count=len(entries), entries=entries)
    with open(args.out, "w", encoding="utf-8") as fp:
        json.dump(payload, fp, ensure_ascii=False, indent=1)
    pref = len({e["pref"] for e in entries})
    print(f"書き出しました: {args.out}")
    print(f"  {len(entries)} 件 / {pref} 都道府県（出典: {source}）")


if __name__ == "__main__":
    main()
